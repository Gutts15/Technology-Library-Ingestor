#!/usr/bin/env python3
"""Selective local spreadsheet ingestion for Technology Library.

V1 supports CSV, TSV, XLSX and XLSM. It preserves workbook structure, source row
numbers, formulas, and cached formula values when the workbook provides them.
Data is emitted as sparse JSONL chunks so future retrieval can load only a
specific sheet/range instead of the full workbook.

No external API or semantic LLM analysis is performed. Normal stdout contains
only aggregate counts, never private filenames, sheet names or cell contents.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

SPREADSHEET_PIPELINE_VERSION = "0.1.0"
SPREADSHEET_SUMMARY_VERSION = "0.1.0"
SCHEMA_VERSION = 1
MAX_SUMMARY_BYTES = 3072
MAX_INPUT_BYTES = 250 * 1024 * 1024
DEFAULT_ROWS_PER_CHUNK = 1000
MAX_SAMPLE_CELL_CHARS = 180
MAX_SAMPLE_COLUMNS = 12

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm"}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def json_value(value: Any) -> Any:
    """Convert common spreadsheet values to deterministic JSON-compatible data."""
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def compact_cell_value(value: Any, limit: int = MAX_SAMPLE_CELL_CHARS) -> str:
    if value is None:
        return ""
    rendered = str(json_value(value)).replace("\r", " ").replace("\n", " ").strip()
    return rendered[:limit]


def encoded_size(payload: dict[str, Any]) -> int:
    return len(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def compact_summary(payload: dict[str, Any]) -> dict[str, Any]:
    if encoded_size(payload) <= MAX_SUMMARY_BYTES:
        return payload

    samples = payload.get("evidence", {}).get("row_samples") if isinstance(payload.get("evidence"), dict) else None
    if isinstance(samples, list):
        for cell_limit in (100, 60, 30):
            payload["evidence"]["row_samples"] = [
                [str(cell)[:cell_limit] for cell in row[:8]] for row in samples[:2]
            ]
            if encoded_size(payload) <= MAX_SUMMARY_BYTES:
                return payload
        payload["evidence"]["row_samples"] = []

    if encoded_size(payload) > MAX_SUMMARY_BYTES:
        payload.pop("evidence", None)
    if encoded_size(payload) > MAX_SUMMARY_BYTES:
        payload = {
            "schema_version": payload.get("schema_version"),
            "summary_version": payload.get("summary_version"),
            "source": payload.get("source"),
            "spreadsheet": payload.get("spreadsheet"),
            "audit": payload.get("audit"),
        }
    return payload


def chunk_path_for(sheet_id: str, chunk_index: int) -> str:
    return f"sheets/{sheet_id}/chunk-{chunk_index:04d}.jsonl"


def write_jsonl_chunk(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")


def choose_header_candidate(nonempty_rows: list[tuple[int, list[Any]]]) -> tuple[int | None, list[str]]:
    if not nonempty_rows:
        return None, []
    row_number, values = nonempty_rows[0]
    headers = [compact_cell_value(value, 120) for value in values[:50]]
    while headers and not headers[-1]:
        headers.pop()
    return row_number, headers


def row_sample(values: list[Any]) -> list[str]:
    return [compact_cell_value(value) for value in values[:MAX_SAMPLE_COLUMNS]]


def sparse_row_from_values(row_number: int, values: list[Any]) -> dict[str, Any] | None:
    cells: list[dict[str, Any]] = []
    for column, value in enumerate(values, start=1):
        if value is None or value == "":
            continue
        cells.append({"c": column, "v": json_value(value)})
    return {"row": row_number, "cells": cells} if cells else None


def sparse_row_from_excel(
    row_number: int,
    formula_cells: Iterable[Any],
    value_cells: Iterable[Any],
) -> tuple[dict[str, Any] | None, int, int, list[Any]]:
    cells: list[dict[str, Any]] = []
    formula_count = 0
    nonempty = 0
    flat_values: list[Any] = []

    for column, (formula_cell, value_cell) in enumerate(zip(formula_cells, value_cells), start=1):
        formula_value = formula_cell.value
        cached_value = value_cell.value
        display_value = cached_value if cached_value is not None else formula_value
        flat_values.append(display_value)

        is_formula = formula_cell.data_type == "f" or (
            isinstance(formula_value, str) and formula_value.startswith("=")
        )
        if is_formula:
            formula_count += 1
            nonempty += 1
            cell: dict[str, Any] = {"c": column, "f": str(formula_value)}
            if cached_value is not None:
                cell["cached"] = json_value(cached_value)
            cells.append(cell)
            continue

        if formula_value is None or formula_value == "":
            continue
        nonempty += 1
        cells.append({"c": column, "v": json_value(formula_value)})

    return ({"row": row_number, "cells": cells} if cells else None), nonempty, formula_count, flat_values


def emit_sheet_chunks(
    *,
    out: Path,
    sheet_id: str,
    sparse_rows: Iterable[dict[str, Any]],
    rows_per_chunk: int,
) -> tuple[list[dict[str, Any]], int]:
    chunks: list[dict[str, Any]] = []
    buffer: list[dict[str, Any]] = []
    total_nonempty_rows = 0

    def flush() -> None:
        nonlocal buffer
        if not buffer:
            return
        chunk_index = len(chunks) + 1
        relative = chunk_path_for(sheet_id, chunk_index)
        write_jsonl_chunk(out / relative, buffer)
        chunks.append({
            "index": chunk_index,
            "path": relative,
            "start_row": buffer[0]["row"],
            "end_row": buffer[-1]["row"],
            "rows": len(buffer),
            "cells": sum(len(row["cells"]) for row in buffer),
        })
        buffer = []

    for row in sparse_rows:
        total_nonempty_rows += 1
        buffer.append(row)
        if len(buffer) >= rows_per_chunk:
            flush()
    flush()
    return chunks, total_nonempty_rows


def parse_delimited(path: Path, out: Path, delimiter: str, rows_per_chunk: int) -> tuple[list[dict[str, Any]], dict[str, int], list[list[str]]]:
    sheet_id = "sheet-0001"
    chunks_buffer: list[dict[str, Any]] = []
    chunks: list[dict[str, Any]] = []
    nonempty_rows: list[tuple[int, list[Any]]] = []
    samples: list[list[str]] = []
    max_columns = 0
    nonempty_cells = 0
    last_nonempty_row = 0

    def flush() -> None:
        nonlocal chunks_buffer
        if not chunks_buffer:
            return
        idx = len(chunks) + 1
        relative = chunk_path_for(sheet_id, idx)
        write_jsonl_chunk(out / relative, chunks_buffer)
        chunks.append({
            "index": idx,
            "path": relative,
            "start_row": chunks_buffer[0]["row"],
            "end_row": chunks_buffer[-1]["row"],
            "rows": len(chunks_buffer),
            "cells": sum(len(row["cells"]) for row in chunks_buffer),
        })
        chunks_buffer = []

    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row_number, row in enumerate(reader, start=1):
            max_columns = max(max_columns, len(row))
            sparse = sparse_row_from_values(row_number, row)
            if sparse is None:
                continue
            last_nonempty_row = row_number
            nonempty_cells += len(sparse["cells"])
            if len(nonempty_rows) < 1:
                nonempty_rows.append((row_number, row))
            if len(samples) < 3:
                samples.append(row_sample(row))
            chunks_buffer.append(sparse)
            if len(chunks_buffer) >= rows_per_chunk:
                flush()
    flush()

    header_row, headers = choose_header_candidate(nonempty_rows)
    sheet = {
        "id": sheet_id,
        "name": "CSV" if delimiter == "," else "TSV",
        "index": 1,
        "state": "visible",
        "last_nonempty_row": last_nonempty_row,
        "max_column": max_columns,
        "nonempty_rows": sum(chunk["rows"] for chunk in chunks),
        "nonempty_cells": nonempty_cells,
        "formula_cells": 0,
        "header_candidate_row": header_row,
        "header_candidate": headers,
        "chunks": chunks,
    }
    totals = {
        "rows": sheet["nonempty_rows"],
        "cells": nonempty_cells,
        "formulas": 0,
        "chunks": len(chunks),
    }
    return [sheet], totals, samples


def load_excel_workbooks(path: Path) -> tuple[Any, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl_missing") from exc

    try:
        formulas = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        values = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        return formulas, values
    except Exception as exc:
        raise RuntimeError("workbook_open_failed") from exc


def parse_excel(path: Path, out: Path, rows_per_chunk: int) -> tuple[list[dict[str, Any]], dict[str, int], list[list[str]]]:
    formulas_wb, values_wb = load_excel_workbooks(path)
    sheets_index: list[dict[str, Any]] = []
    samples: list[list[str]] = []
    totals = {"rows": 0, "cells": 0, "formulas": 0, "chunks": 0}

    try:
        for sheet_number, formula_ws in enumerate(formulas_wb.worksheets, start=1):
            sheet_id = f"sheet-{sheet_number:04d}"
            try:
                value_ws = values_wb[formula_ws.title]
            except KeyError as exc:
                raise RuntimeError("sheet_alignment_failed") from exc

            chunks_buffer: list[dict[str, Any]] = []
            chunks: list[dict[str, Any]] = []
            first_nonempty: tuple[int, list[Any]] | None = None
            last_nonempty_row = 0
            max_column = 0
            nonempty_cells = 0
            formula_cells = 0
            nonempty_rows_count = 0

            def flush() -> None:
                nonlocal chunks_buffer
                if not chunks_buffer:
                    return
                idx = len(chunks) + 1
                relative = chunk_path_for(sheet_id, idx)
                write_jsonl_chunk(out / relative, chunks_buffer)
                chunks.append({
                    "index": idx,
                    "path": relative,
                    "start_row": chunks_buffer[0]["row"],
                    "end_row": chunks_buffer[-1]["row"],
                    "rows": len(chunks_buffer),
                    "cells": sum(len(row["cells"]) for row in chunks_buffer),
                })
                chunks_buffer = []

            formula_iter = formula_ws.iter_rows()
            value_iter = value_ws.iter_rows()
            for row_number, (formula_row, value_row) in enumerate(zip(formula_iter, value_iter), start=1):
                sparse, row_nonempty, row_formulas, flat_values = sparse_row_from_excel(
                    row_number, formula_row, value_row
                )
                if sparse is None:
                    continue
                nonempty_rows_count += 1
                last_nonempty_row = row_number
                max_column = max(max_column, max(cell["c"] for cell in sparse["cells"]))
                nonempty_cells += row_nonempty
                formula_cells += row_formulas
                if first_nonempty is None:
                    first_nonempty = (row_number, flat_values)
                if len(samples) < 3:
                    samples.append(row_sample(flat_values))
                chunks_buffer.append(sparse)
                if len(chunks_buffer) >= rows_per_chunk:
                    flush()
            flush()

            header_row, headers = choose_header_candidate([first_nonempty] if first_nonempty else [])
            sheet_entry = {
                "id": sheet_id,
                "name": formula_ws.title,
                "index": sheet_number,
                "state": formula_ws.sheet_state,
                "last_nonempty_row": last_nonempty_row,
                "max_column": max_column,
                "nonempty_rows": nonempty_rows_count,
                "nonempty_cells": nonempty_cells,
                "formula_cells": formula_cells,
                "header_candidate_row": header_row,
                "header_candidate": headers,
                "chunks": chunks,
            }
            sheets_index.append(sheet_entry)
            totals["rows"] += nonempty_rows_count
            totals["cells"] += nonempty_cells
            totals["formulas"] += formula_cells
            totals["chunks"] += len(chunks)
    finally:
        formulas_wb.close()
        values_wb.close()

    return sheets_index, totals, samples


def build_summary(
    *,
    source_id: str | None,
    source_sha: str,
    source_type: str,
    sheets: list[dict[str, Any]],
    totals: dict[str, int],
    samples: list[list[str]],
) -> dict[str, Any]:
    payload = {
        "schema_version": SCHEMA_VERSION,
        "summary_version": SPREADSHEET_SUMMARY_VERSION,
        "source": {
            "provider": "external",
            "file_id": source_id,
            "sha256": source_sha,
        },
        "spreadsheet": {
            "type": source_type,
            "sheets": len(sheets),
            "nonempty_rows": totals["rows"],
            "nonempty_cells": totals["cells"],
            "formula_cells": totals["formulas"],
            "chunks": totals["chunks"],
        },
        "evidence": {"row_samples": samples[:3]},
        "audit": {
            "manifest": "ingest.json",
            "index": "spreadsheet-index.json",
            "chunks": "sheets/",
        },
    }
    return compact_summary(payload)


def main() -> int:
    parser = argparse.ArgumentParser(description="Ingest CSV/TSV/XLSX/XLSM into selective spreadsheet evidence.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--source-id", default=None)
    parser.add_argument("--rows-per-chunk", type=int, default=DEFAULT_ROWS_PER_CHUNK)
    args = parser.parse_args()

    source = args.source.resolve()
    out = args.out.resolve()
    if not source.is_file():
        print("spreadsheet_ingest_error code=source_missing")
        return 2
    if source.stat().st_size > MAX_INPUT_BYTES:
        print("spreadsheet_ingest_error code=input_too_large")
        return 2
    if not 100 <= args.rows_per_chunk <= 10000:
        parser.error("--rows-per-chunk must be between 100 and 10000")

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        print("spreadsheet_ingest_error code=unsupported_format")
        return 3

    source_sha = sha256_file(source)
    out.mkdir(parents=True, exist_ok=True)

    try:
        if suffix == ".csv":
            sheets, totals, samples = parse_delimited(source, out, ",", args.rows_per_chunk)
            source_type = "csv"
        elif suffix == ".tsv":
            sheets, totals, samples = parse_delimited(source, out, "\t", args.rows_per_chunk)
            source_type = "tsv"
        else:
            sheets, totals, samples = parse_excel(source, out, args.rows_per_chunk)
            source_type = "xlsm" if suffix == ".xlsm" else "xlsx"
    except RuntimeError as exc:
        print(f"spreadsheet_ingest_error code={exc}")
        return 4
    except Exception:
        print("spreadsheet_ingest_error code=spreadsheet_parse_failed")
        return 4

    index_payload = {
        "schema_version": SCHEMA_VERSION,
        "pipeline_version": SPREADSHEET_PIPELINE_VERSION,
        "source_type": source_type,
        "sheets": sheets,
        "totals": totals,
    }
    write_json(out / "spreadsheet-index.json", index_payload)

    manifest = {
        "schema_version": SCHEMA_VERSION,
        "pipeline_version": SPREADSHEET_PIPELINE_VERSION,
        "processing": {"status": "processed", "processed_at": utc_now()},
        "source": {
            "provider": "external",
            "file_id": args.source_id,
            "sha256": source_sha,
            "size_bytes": source.stat().st_size,
        },
        "spreadsheet": {
            "type": source_type,
            "sheets": len(sheets),
            "nonempty_rows": totals["rows"],
            "nonempty_cells": totals["cells"],
            "formula_cells": totals["formulas"],
            "chunks": totals["chunks"],
            "macros_extracted": False,
        },
    }
    write_json(out / "ingest.json", manifest)

    summary = build_summary(
        source_id=args.source_id,
        source_sha=source_sha,
        source_type=source_type,
        sheets=sheets,
        totals=totals,
        samples=samples,
    )
    summary_path = out / "evidence-summary.json"
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    if summary_path.stat().st_size > MAX_SUMMARY_BYTES:
        print("spreadsheet_ingest_error code=summary_budget_exceeded")
        return 5

    write_json(out / "checkpoint.json", {
        "schema_version": SCHEMA_VERSION,
        "status": "processed",
        "source_sha256": source_sha,
        "pipeline_version": SPREADSHEET_PIPELINE_VERSION,
        "summary_version": SPREADSHEET_SUMMARY_VERSION,
        "updated_at": utc_now(),
    })

    print(
        "spreadsheet_ingest_ok "
        f"type={source_type} sheets={len(sheets)} rows={totals['rows']} "
        f"cells={totals['cells']} formulas={totals['formulas']} chunks={totals['chunks']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
